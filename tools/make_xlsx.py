import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# 사용법: python3 tools/make_xlsx.py  (저장소 루트에서 실행)
import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
args = sys.argv[1:] + [None] * 3
MD_PATH = args[0] or os.path.join(ROOT, '아파트_셀프_사전점검_체크리스트.md')
out = args[1] or os.path.join(ROOT, '아파트_셀프_사전점검_체크리스트.xlsx')
STD_JSON = args[2] or os.path.join(ROOT, 'data', 'standard_form.json')
VIDEOS_TSV = os.path.join(ROOT, 'data', 'youtube_videos.tsv')
md = open(MD_PATH, encoding='utf-8').read()

def section(title_prefix):
    m = re.search(r'^## ' + re.escape(title_prefix) + r'.*?$(.*?)(?=^## |\Z)', md, re.M | re.S)
    return m.group(1)

def clean(t):
    t = re.sub(r'\*\*(.+?)\*\*', r'\1', t)
    t = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', t)
    t = re.sub(r'\*\((.+?)\)\*', r'(\1)', t)
    return t.strip()

# prep items
prep = [clean(l[6:]) for l in section('3.').splitlines() if l.startswith('- [ ] ')]
# inspection items by space
rooms = []
cur = None; sub = ''
for l in section('4.').splitlines():
    if l.startswith('### '):
        name = re.sub(r'^[^\w가-힣]+', '', l[4:].replace('*', '')).strip()
        name = re.sub(r'\s*\(.*\)$', '', name)
        cur = [name, []]; rooms.append(cur); sub = ''
    elif l.startswith('**') and l.rstrip().endswith('**'):
        sub = l.strip('* \n')
    elif l.startswith('- [ ] ') and cur:
        cur[1].append((sub, clean(l[6:])))

FONT = '맑은 고딕'
f = lambda **k: Font(name=FONT, **k)
thin = Side(style='thin', color='B7C4CF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='2F5D7C')
ROOMFILL = PatternFill('solid', fgColor='DCE8F0')
INPUT = PatternFill('solid', fgColor='FFF9DB')
wrap = Alignment(wrap_text=True, vertical='center')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = f(bold=True, color='FFFFFF'); cell.fill = HDR
        cell.alignment = center; cell.border = border
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[row].height = 24

wb = Workbook()

# ---- 안내
ws = wb.active; ws.title = '안내'
ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 80
ws['A1'] = '아파트 셀프 사전점검 체크리스트'; ws['A1'].font = f(bold=True, size=16, color='2F5D7C')
ws['A2'] = '작성일 2026-10-08 · YouTube 조회수 상위 사전점검 영상 9편 자막 + 국토교통부 표준점검표 기준'
ws['A2'].font = f(size=10, color='555555')
rows = [
    ('시트 구성', ''),
    ('준비물', '가기 전에 챙길 것. 챙기면 B열에 ○ 선택'),
    ('점검표', '공간별 셀프 점검 항목(영상·가이드 기반 상세판). 노란 칸만 입력하세요'),
    ('표준점검표', '국토교통부 입주자 사전방문 표준점검표(포스코이앤씨 배포본) 136항목. 제출용 점검표와 같은 구성'),
    ('하자기록', '하자를 찾으면 한 줄씩 기록. K열 문구를 복사해 더샵 앱에 붙여넣기'),
    ('더샵 사전방문', '포스코이앤씨 더샵 사전방문 예약부터 휴대폰 접수, 보수 확인까지 순서'),
    ('요약', '공간별 진행률과 하자 개수가 자동 계산됩니다 (입력 불필요)'),
    ('참고영상', '참고한 YouTube 영상과 자료'),
    ('', ''),
    ('결과 입력 방법', ''),
    ('○', '이상 없음'),
    ('✖', '하자 있음 → 하자번호를 적고 하자기록 시트에 상세 기록'),
    ('△', '애매함 → 나중에 다시 확인'),
    ('', ''),
    ('점검 순서', '현관에서 시작해 시계 방향으로. 공간마다 문 → 천장 → 창 → 벽 → 바닥 → 작동 테스트'),
    ('사진 원칙', '원거리 1장(위치) + 근접 1장(하자). 배수·소음은 동영상'),
    ('접수', '당일 시공사 앱으로 접수 후 화면 캡처. 앱 오류 시 현장 본부에 서면 하자보수 요청서 제출'),
    ('셀프 한계', '열화상(단열·결로), 라돈·공기질, 콘센트 접지, 매립 배관 누수는 장비 필요'),
    ('', ''),
    ('출처', '점검표 항목 중 (영상) 표시는 YouTube 영상 자막에서, (표준점검표) 표시는 국토교통부 표준점검표에서 가져온 것입니다.'),
]
for i, (a, b) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=a).font = f(bold=bool(a) and not b or a in ('주의',), size=11)
    c = ws.cell(row=i, column=2, value=b); c.font = f(size=11); c.alignment = wrap
    if a and not b:
        ws.cell(row=i, column=1).font = f(bold=True, size=12, color='2F5D7C')


res_dv = lambda: DataValidation(type='list', formula1='"○,✖,△"', allow_blank=True)

# ---- 준비물
ws = wb.create_sheet('준비물')
header(ws, 1, ['준비물', '챙김', '메모'], [60, 8, 30])
dv = DataValidation(type='list', formula1='"○"', allow_blank=True); ws.add_data_validation(dv)
for i, p in enumerate(prep, 2):
    ws.cell(row=i, column=1, value=p).font = f(); ws.cell(row=i, column=1).alignment = wrap
    for c in (2, 3):
        cell = ws.cell(row=i, column=c); cell.fill = INPUT; cell.font = f(); cell.border = border
        cell.alignment = center if c == 2 else wrap
    ws.cell(row=i, column=1).border = border
    dv.add(f'B{i}')
last = len(prep) + 1
ws.cell(row=last + 2, column=1, value='챙긴 개수').font = f(bold=True)
ws.cell(row=last + 2, column=2, value=f'=COUNTIF(B2:B{last},"○")&"/"&COUNTA(A2:A{last})').font = f(bold=True)
ws.freeze_panes = 'A2'

# ---- 점검표
ws = wb.create_sheet('점검표')
cols = ['No', '공간', '구분', '점검 항목', '결과', '하자번호', '위치 / 메모']
header(ws, 1, cols, [6, 16, 10, 62, 8, 10, 34])
dv = res_dv(); ws.add_data_validation(dv)
r = 2; n = 0
for name, items in rooms:
    for sub, text in items:
        n += 1
        vals = [n, name, sub, text]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = f(); cell.border = border
            cell.alignment = center if c in (1, 3) else wrap
        for c in (5, 6, 7):
            cell = ws.cell(row=r, column=c); cell.fill = INPUT; cell.border = border; cell.font = f()
            cell.alignment = center if c < 7 else wrap
        dv.add(f'E{r}')
        r += 1
LAST = r - 1
ws.freeze_panes = 'E2'
ws.auto_filter.ref = f'A1:G{LAST}'
red = PatternFill('solid', fgColor='F8D7DA'); green = PatternFill('solid', fgColor='D4EDDA'); amber = PatternFill('solid', fgColor='FFE8B3')
for sym, fill in (('✖', red), ('○', green), ('△', amber)):
    ws.conditional_formatting.add(f'E2:E{LAST}', CellIsRule(operator='equal', formula=[f'"{sym}"'], fill=fill))
ws.print_title_rows = '1:1'
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---- 하자기록
ws = wb.create_sheet('하자기록')
cols = ['하자번호', '공간', '위치(구체적으로)', '증상', '사진', '접수일', '접수 방법', '보수 예정일', '보수 완료 확인', '비고', '앱 접수 문구 (자동 · 복사해서 붙여넣기)']
header(ws, 1, cols, [9, 14, 30, 34, 8, 12, 12, 12, 12, 24, 48])
space_names = [nm for nm, _ in rooms]
wb.create_sheet('_목록').sheet_state = 'hidden'
lst = wb['_목록']
for i, nm in enumerate(space_names, 1): lst.cell(row=i, column=1, value=nm)
dv_sp = DataValidation(type='list', formula1=f"='_목록'!$A$1:$A${len(space_names)}", allow_blank=True)
dv_ph = DataValidation(type='list', formula1='"○,×"', allow_blank=True)
dv_how = DataValidation(type='list', formula1='"더샵 앱/웹,서면,현장"', allow_blank=True)
dv_done = DataValidation(type='list', formula1='"완료,미완료,재요청"', allow_blank=True)
dv_date = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True)
for d in (dv_sp, dv_ph, dv_how, dv_done, dv_date): ws.add_data_validation(d)
example = ['예시) 1', space_names[3] if len(space_names) > 3 else '', '안방 욕실 바닥, 배수구 왼쪽 30cm', '물 부으면 배수구 반대편에 고임 (영상 촬영)', '○', None, '더샵 앱/웹', None, '미완료', '이 줄은 예시입니다. 지우고 쓰세요']
import datetime
example[5] = datetime.date(2026, 10, 8); example[7] = datetime.date(2026, 11, 15)
for rr in range(2, 102):
    for c in range(1, 11):
        cell = ws.cell(row=rr, column=c); cell.border = border; cell.font = f(); cell.fill = INPUT
        cell.alignment = wrap if c in (3, 4, 10) else center
        if c in (6, 8): cell.number_format = 'yyyy-mm-dd'
    if rr > 2: ws.cell(row=rr, column=1, value=rr - 2)
    k = ws.cell(row=rr, column=11, value=f'=IF(C{rr}="","",B{rr}&" / "&C{rr}&" : "&D{rr})')
    k.border = border; k.font = f(); k.alignment = wrap
    dv_sp.add(f'B{rr}'); dv_ph.add(f'E{rr}'); dv_how.add(f'G{rr}'); dv_done.add(f'I{rr}'); dv_date.add(f'F{rr}'); dv_date.add(f'H{rr}')
for c, v in enumerate(example, 1):
    cell = ws.cell(row=2, column=c, value=v); cell.font = f(italic=True, color='7F7F7F'); cell.fill = PatternFill('solid', fgColor='F2F2F2')
ws.conditional_formatting.add('I3:I101', CellIsRule(operator='equal', formula=['"완료"'], fill=green))
ws.conditional_formatting.add('I3:I101', CellIsRule(operator='equal', formula=['"재요청"'], fill=red))
ws.freeze_panes = 'B2'; ws.auto_filter.ref = 'A1:K101'

# ---- 요약
ws = wb.create_sheet('요약')
header(ws, 1, ['공간', '항목 수', '점검 완료', '진행률', '○ 이상없음', '✖ 하자', '△ 재확인', '하자기록 수', '보수 완료'], [18, 9, 10, 9, 11, 9, 10, 11, 10])
P = "점검표"
for i, nm in enumerate(space_names, 2):
    ws.cell(row=i, column=1, value=nm)
    ws.cell(row=i, column=2, value=f'=COUNTIF({P}!$B$2:$B${LAST},A{i})')
    ws.cell(row=i, column=3, value=f'=COUNTIFS({P}!$B$2:$B${LAST},A{i},{P}!$E$2:$E${LAST},"<>")')
    ws.cell(row=i, column=4, value=f'=IF(B{i}=0,0,C{i}/B{i})')
    for c, s in ((5, '○'), (6, '✖'), (7, '△')):
        ws.cell(row=i, column=c, value=f'=COUNTIFS({P}!$B$2:$B${LAST},A{i},{P}!$E$2:$E${LAST},"{s}")')
    ws.cell(row=i, column=8, value=f'=COUNTIF(하자기록!$B$3:$B$101,A{i})')
    ws.cell(row=i, column=9, value=f'=COUNTIFS(하자기록!$B$3:$B$101,A{i},하자기록!$I$3:$I$101,"완료")')
t = len(space_names) + 2
ws.cell(row=t, column=1, value='합계')
for c in range(2, 10):
    L = ws.cell(row=1, column=c).column_letter
    ws.cell(row=t, column=c, value=f'=IF(B{t}=0,0,C{t}/B{t})' if c == 4 else f'=SUM({L}2:{L}{t-1})')
for rr in range(2, t + 1):
    for c in range(1, 10):
        cell = ws.cell(row=rr, column=c); cell.border = border; cell.font = f(bold=(rr == t))
        cell.alignment = center if c > 1 else wrap
        if c == 4: cell.number_format = '0%'
        if rr == t: cell.fill = ROOMFILL
ws.cell(row=t + 2, column=1, value='※ 하자기록 수는 2행 예시를 제외하고 셉니다.').font = f(size=9, color='7F7F7F')
from openpyxl.formatting.rule import DataBarRule
ws.conditional_formatting.add(f'D2:D{t-1}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='5B9BD5'))


# ---- 표준점검표 (국토교통부 입주자 사전방문 표준점검표, 포스코이앤씨 배포본)
import json
std_rows = json.load(open(STD_JSON, encoding='utf-8'))
ws = wb.create_sheet('표준점검표')
ws['A1'] = '입주자 사전방문 표준점검표'; ws['A1'].font = f(bold=True, size=14, color='2F5D7C')
ws['A2'] = '포스코이앤씨 모바일 배포 PDF(국토교통부 표준 양식)를 옮긴 것입니다. 거실·침실·욕실이 여러 개면 위치/메모에 번호(침실1, 욕실2 등)를 적으세요.'; ws['A2'].font = f(size=10, color='555555')
for i, lab in enumerate(['동호수', '점검일', '점검자', '연락처']):
    c = 1 + i * 2
    a = ws.cell(row=3, column=c, value=lab); a.font = f(bold=True); a.border = border; a.alignment = center
    b = ws.cell(row=3, column=c + 1); b.fill = INPUT; b.border = border; b.font = f()
ws['D3'].number_format = 'yyyy-mm-dd'
header(ws, 5, ['No', '구역', '부위', '점검대상', '점검사항', '결과', '조치요구사항', '하자번호'], [6, 18, 7, 30, 32, 8, 40, 9])
dvs = res_dv(); ws.add_data_validation(dvs)
for i, (sec, part, tgt, chk) in enumerate(std_rows, 6):
    for c, v in enumerate([i - 5, sec, part, tgt, chk], 1):
        x = ws.cell(row=i, column=c, value=v); x.font = f(); x.border = border; x.alignment = center if c in (1, 3) else wrap
    for c in (6, 7, 8):
        x = ws.cell(row=i, column=c); x.fill = INPUT; x.border = border; x.font = f(); x.alignment = wrap if c == 7 else center
    dvs.add(f'F{i}')
SL = 5 + len(std_rows)
for sym, fill in (('✖', red), ('○', green), ('△', amber)):
    ws.conditional_formatting.add(f'F6:F{SL}', CellIsRule(operator='equal', formula=[f'"{sym}"'], fill=fill))
ws.freeze_panes = 'F6'; ws.auto_filter.ref = f'A5:H{SL}'
ws.cell(row=SL + 2, column=2, value='점검 완료').font = f(bold=True)
ws.cell(row=SL + 2, column=4, value=f'=COUNTA(F6:F{SL})&" / "&COUNTA(E6:E{SL})&" 항목"').font = f(bold=True)
ws.cell(row=SL + 3, column=2, value='하자(✖)').font = f(bold=True)
ws.cell(row=SL + 3, column=4, value=f'=COUNTIF(F6:F{SL},"✖")').font = f(bold=True, color='C0392B')
ws.print_title_rows = '5:5'; ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---- 더샵 사전방문
ws = wb.create_sheet('더샵 사전방문')
ws.column_dimensions['A'].width = 6; ws.column_dimensions['B'].width = 64; ws.column_dimensions['C'].width = 8; ws.column_dimensions['D'].width = 36
ws['A1'] = '포스코이앤씨 더샵 사전방문(사전점검) · 휴대폰 하자 접수 순서'; ws['A1'].font = f(bold=True, size=14, color='2F5D7C')
ws['A2'] = '※ 단지마다 앱 이름·접수 방식·기한이 다릅니다. 노란 칸은 받으신 사전방문 안내문(문자·우편)을 보고 채우세요.'; ws['A2'].font = f(size=10, color='C0392B')
info = ['단지명', '사전방문 일시 (예약한 날짜·시간)', '하자 접수 앱/웹 이름·주소 (안내문 기재)', '앱 로그인 정보 (계약자명·동호수)', '하자 접수 마감 일시', '입주지원센터 연락처', '보수 확인 재방문 일정']
ws['A4'] = '기본 정보'; ws['A4'].font = f(bold=True, size=12, color='2F5D7C')
for i, t in enumerate(info, 5):
    a = ws.cell(row=i, column=2, value=t); a.font = f(); a.border = border
    for c in (3, 4):
        b = ws.cell(row=i, column=c); b.fill = INPUT; b.border = border; b.font = f()
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=4)
r0 = 5 + len(info) + 1
ws.cell(row=r0, column=1, value='진행 순서').font = f(bold=True, size=12, color='2F5D7C')
for c, h in enumerate(['단계', '할 일', '완료', '메모'], 1):
    x = ws.cell(row=r0 + 1, column=c, value=h); x.font = f(bold=True, color='FFFFFF'); x.fill = HDR; x.alignment = center; x.border = border
steps = [
 ('사전', '안내문(문자·우편)에서 방문 예약 → 날짜·시간 확정'),
 ('사전', '단지 지정 하자 접수 앱/웹을 미리 설치하고 로그인 확인 (계약자 정보·동호수 입력)'),
 ('사전', '포스코이앤씨 모바일의 입주자 사전방문 표준점검표 PDF 확인 → 엑셀 표준점검표 시트와 같은 구성'),
 ('사전', '방문 전날 받은 입장 QR코드(단지에 따라 발송) 캡처해 두기'),
 ('사전', '휴대폰 충전, 보조배터리, 저장공간 확보 (사진이 많음)'),
 ('당일', '입주지원센터(접수처)에서 QR·신분증 확인 → 교육 듣기 → 점검표·스티커 받기'),
 ('당일', '현관부터 시계 방향으로 점검 (점검표 시트 순서대로)'),
 ('당일', '하자 발견 → 스티커 부착 → 사진(원거리+근접) → 하자기록 시트에 기록'),
 ('당일', '휴대폰 앱/웹에 하자 등록: 공간 선택 → 사진 첨부 → 위치·증상 입력 (하자기록 K열 문구 복사)'),
 ('당일', '나가기 전 접수 목록 화면을 캡처하고, 하자기록과 개수 맞춰 보기'),
 ('기한 내', '집에서 사진을 다시 보며 빠진 하자 추가 등록 (접수 마감 전까지)'),
 ('이후', '앱에서 처리 상태(접수 → 보수 중 → 완료) 확인'),
 ('이후', '보수 확인 재방문 때 하자번호별로 완료 확인 → 하자기록 I열 갱신'),
 ('이후', '앱에 안 되는 건은 입주지원센터에 서면 요청서 제출, 입주 후에는 관리사무소·하자보수 기간 내 접수'),
]
dvd = DataValidation(type='list', formula1='"○"', allow_blank=True); ws.add_data_validation(dvd)
for i, (st, t) in enumerate(steps, r0 + 2):
    for c, v in ((1, st), (2, t)):
        x = ws.cell(row=i, column=c, value=v); x.font = f(); x.border = border; x.alignment = center if c == 1 else wrap
    for c in (3, 4):
        x = ws.cell(row=i, column=c); x.fill = INPUT; x.border = border; x.font = f(); x.alignment = center if c == 3 else wrap
    dvd.add(f'C{i}')
r1 = r0 + 2 + len(steps) + 1
ws.cell(row=r1, column=1, value='하자로 잘못 접수하기 쉬운 것 (현장 확인 후 판단)').font = f(bold=True, size=12, color='2F5D7C')
notdef = ['현관 방화문에 말발굽(도어스토퍼)이 없는 것 — 방화문은 닫힌 상태가 원칙이라 기본 미설치인 경우가 많음',
          '목문틀의 타카핀 자국 — 시공 방식상 생기는 자국',
          '줄눈 없는(무메지) 아트월 타일 — 설계상 줄눈이 없는 마감',
          '계약하지 않은 옵션이 없는 것 — 분양·옵션 계약서와 대조']
for i, t in enumerate(notdef, r1 + 1):
    x = ws.cell(row=i, column=2, value='· ' + t); x.font = f(); x.alignment = wrap
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    ws.row_dimensions[i].height = 30
ws.cell(row=r1 + len(notdef) + 2, column=2, value='출처: 광명 자이 더샵 포레나 사전점검 가이드(xi-guide.oopy.io), 사전점검 앱 공통 접수 흐름 자료. 더샵 공식 절차는 단지 안내문이 우선입니다.').font = f(size=9, color='7F7F7F')

# ---- 참고영상
ws = wb.create_sheet('참고영상')
header(ws, 1, ['#', '영상', '채널', '조회수', '좋아요', '올린 날', '링크'], [5, 52, 20, 11, 9, 11, 40])
import csv
with open(VIDEOS_TSV, encoding='utf-8') as fh:
    vids = list(csv.DictReader(fh, delimiter='\t'))
for i, v in enumerate(vids, 2):
    d = v.get('upload_date') or ''
    vals = [i - 1, v['title'], v['channel'],
            int(v['views']) if v['views'].isdigit() else '-',
            int(v['likes']) if v['likes'].isdigit() else '-',
            f'{d[:4]}-{d[4:6]}-{d[6:]}' if len(d) == 8 else '',
            f"https://www.youtube.com/watch?v={v['id']}"]
    for c, val in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=val); cell.font = f(); cell.border = border
        cell.alignment = wrap if c in (2, 3, 7) else center
        if c in (4, 5): cell.number_format = '#,##0'
    ws.cell(row=i, column=7).hyperlink = vals[6]
ws.cell(row=len(vids) + 3, column=2, value='※ 조회수 상위 해설 영상의 자막 전체를 읽고 반영했습니다(2026-10-08). 새 영상은 tools/yt_tool.py로 추가합니다.').font = f(size=9, color='7F7F7F')

wb.move_sheet('_목록', offset=10)
wb.calculation.fullCalcOnLoad = True
wb.save(out)
print(len(prep), 'prep;', LAST - 1, 'items;', [ (nm, len(it)) for nm, it in rooms])
